#!/usr/bin/env python3
"""xl_utils — Excel 处理共享工具（skills/xlsx/scripts 下 CLI 的公共底座）

定位：把 xl-read / xl-clean / xl-transform / xl-convert 等技能模板里反复出现的
读表、sheet 维度、列统计、CSV 编码兜底、markdown 生成、JSON 输出等骨架代码
下沉到这里，各 CLI 只做参数编排，不再每轮由 AI 从零重写。

复用约定：
- 输出 JSON 一律走 dump_json()（utf-8、ensure_ascii=False），供宿主稳定解析；
- CSV 读取统一 read_csv_safe()（utf-8-sig → gbk → gb18030 三级兜底）；
- 行数口径：Excel 用 openpyxl read_only 的 max_row（不整包解析），
  大文件（> 50_000 行）只对前 1_000 行做统计并显式标注 sampled。
"""
import json
import os
import re
import sys
import zipfile

import pandas as pd
from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

SAMPLE_LIMIT = 50_000          # 超过则不做全量统计
DEFAULT_SAMPLE_ROWS = 1_000    # 超限时预览/统计的采样行数
CSV_FALLBACK_ENCODINGS = ('utf-8-sig', 'gbk', 'gb18030')
PROBE_MAX_ROWS = 500           # 公式无缓存值探测的行数上限（read_only 流式，够判定了）
CSV_GUARD_WARN_MB = 120        # csv/tsv 全量读入前：超过该体积只警告（stderr，不污染 stdout JSON）
CSV_GUARD_MAX_MB = 400         # csv/tsv 全量读入前：超过该体积直接拒绝（防 pandas 全量 OOM）


def dump_json(obj):
    """统一 JSON 输出（utf-8、中文不转义、datetime/NaN 转字符串）。"""
    print(json.dumps(obj, ensure_ascii=False, default=str, indent=2))


def find_input(path):
    """校验输入存在，给用户可读错误并带非零退出。"""
    if not path or not os.path.isfile(path):
        sys.exit(f'错误：文件不存在 — {path}')
    return path


def sheet_dims(path):
    """openpyxl read_only 模式拿每个 sheet 的真实行/列维度（不整包解析，内存友好）。"""
    dims = {}
    wb = load_workbook(path, read_only=True)
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            dims[name] = (ws.max_row, ws.max_column)
    finally:
        wb.close()
    return dims


def read_csv_safe(path, **kwargs):
    """CSV 编码兜底：utf-8-sig → gbk → gb18030。"""
    last = None
    for enc in CSV_FALLBACK_ENCODINGS:
        try:
            return pd.read_csv(path, encoding=enc, **kwargs)
        except UnicodeDecodeError as e:
            last = e
    raise last


def count_csv_lines(path):
    """流式统计 csv/tsv 的物理行数（内存 O(1)，不解析编码）。

    注意：是近似口径——字段内含引号包裹的换行时物理行数会偏大；
    汇报时建议标注「≈ / 按物理行估算」。
    """
    with open(path, 'rb') as f:
        return sum(1 for _ in f)


def merged_cells_by_sheet(path):
    """用 zipfile 直读各 sheet XML 的 <mergeCell>，统计每 sheet 合并区域数。

    为什么不用 openpyxl：read_only 模式不暴露 merged_cells，普通模式要整包解析
    （与 pandas 全量读同价）。这里只抽取 XML 文本中 mergeCell 标签，成本远低，
    常规 Excel 文件几乎瞬时。非 xlsx/xlsm/xltx 或解析失败返回 {}。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext not in ('.xlsx', '.xlsm', '.xltx'):
        return {}
    try:
        with zipfile.ZipFile(path) as z:
            wbxml = z.read('xl/workbook.xml').decode('utf-8', 'ignore')
            rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8', 'ignore')
    except (KeyError, zipfile.BadZipFile, OSError):
        return {}

    rid2target = {}
    for m in re.finditer(r'<Relationship\b[^>]*>', rels):
        i = re.search(r'Id="(rId\d+)"', m.group(0))
        t = re.search(r'Target="([^"]+)"', m.group(0))
        if i and t:
            rid2target[i.group(1)] = t.group(1)
    sheets = []
    for m in re.finditer(r'<sheet\b[^>]*>', wbxml):
        nm = re.search(r'name="([^"]*)"', m.group(0))
        rid = re.search(r'r:id="(rId\d+)"', m.group(0))
        if nm and rid:
            sheets.append((nm.group(1), rid.group(1)))

    out = {}
    if not sheets:
        return out
    with zipfile.ZipFile(path) as z:
        for name, rid in sheets:
            target = rid2target.get(rid, '')
            if not target:
                continue
            # Target 可能是 'worksheets/sheet1.xml' 或 '/xl/worksheets/sheet1.xml'
            t = target.lstrip('/')
            p = t if t.startswith('xl/') else 'xl/' + t
            if not p.endswith('.xml'):
                continue
            try:
                xml = z.read(p).decode('utf-8', 'ignore')
            except KeyError:
                continue
            out[name] = len(re.findall(r'<mergeCell\b', xml))
    return out


def guard_csv_size(path):
    """csv/tsv 即将被 pandas 全量读入前的体积守卫（只拦整包读，不拦采样）。

    - 超过 CSV_GUARD_MAX_MB：拒绝（pandas 全量解析会膨胀数倍内存，易 OOM/killed）；
    - 超过 CSV_GUARD_WARN_MB：stderr 警告（不污染 stdout 的 JSON 契约），提示走
      采样/分块（见 excel-conventions.md「大文件处理」）。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext not in ('.csv', '.tsv'):
        return
    mb = os.path.getsize(path) / 1e6
    if mb >= CSV_GUARD_MAX_MB:
        sys.exit(f'错误：文件过大（{path} 约 {mb:.0f}MB），pandas 全量读入会占用数倍内存，'
                 f'极易超时/被杀。请先采样确认结构（xl_summarize.py --sample），'
                 f'大文件分批用 read_csv(chunksize=…) 自写脚本处理，不要整包读。')
    if mb >= CSV_GUARD_WARN_MB:
        print(f'警告：{path} 体积约 {mb:.0f}MB，全量读入内存占用较大；'
              f'如超时请改用采样/分块策略。', file=sys.stderr)


def read_frame(path, sheet=None, nrows=None):
    """按文件类型读入 DataFrame：csv/tsv 走编码兜底；xlsx/xlsm 走 read_excel。

    返回 (df, single)：single=True 表示非 Excel 多表文件（无 sheet 概念）。
    注：csv/tsv 全量读（nrows=None）前会做体积守卫；xlsx 加密文件给出可行动报错。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        if nrows is None:
            guard_csv_size(path)
        return read_csv_safe(path, nrows=nrows), True
    if ext == '.tsv':
        if nrows is None:
            guard_csv_size(path)
        return read_csv_safe(path, sep='\t', nrows=nrows), True
    kwargs = {'nrows': nrows} if nrows else {}
    if sheet is not None:
        kwargs['sheet_name'] = sheet
    try:
        return pd.read_excel(path, **kwargs), False
    except Exception as e:
        # 密码保护/加密工作簿：openpyxl 抛的错信息含 encrypted/password 字样，
        # 转成可行动的提示，而不是让 AI 面对晦涩 traceback。
        msg = str(e)
        if re.search(r'encrypt|password', msg, re.I):
            sys.exit(f'错误：{path} 是加密/带密码保护的工作簿，openpyxl 无法直接读取。'
                     f'请让用户在 Excel/WPS 中解密后另存，或按内置 xlsx 技能指引用 '
                     f'msoffcrypto-tool 解密后处理。')
        raise


def probe_formulas(path, sheet=None, max_rows=PROBE_MAX_ROWS):
    """轻量探测 Excel 文件前 max_rows 行是否存在「公式无缓存值」风险。

    背景：openpyxl 写入的公式只有公式串、没有缓存值，pandas（data_only 语义）
    会把这类公式列整列读成 NaN。若宿主据此做空值清洗/删除，会静默破坏公式
    文件。本函数用两趟 read_only 流式读取（公式串 + 缓存值），各至多扫到
    max_rows 行即停，避免整包载入，返回每个 sheet 的探测结果。

    参数：sheet=None 探测全部 sheet；否则只探测指定 sheet。
    返回：{<sheet>: {
            'has_formula': bool,          # 抽查范围内是否存在公式
            'formula_count': int,         # 抽查范围内的公式单元格数
            'no_cache': bool,             # 是否存在公式但缓存值为 None
            'no_cache_count': int,        # 无缓存值的公式单元格数
            'sample_locations': [...]     # 前 5 个无缓存值公式的坐标
          }, ...}；csv/tsv/非 Excel 文件返回 None。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext not in ('.xlsx', '.xlsm', '.xltx'):
        return None

    wb_f = load_workbook(path, read_only=True, data_only=False)
    try:
        wb_v = load_workbook(path, read_only=True, data_only=True)
        try:
            names = [sheet] if sheet else wb_f.sheetnames
            out = {}
            for name in names:
                if name not in wb_f.sheetnames:
                    continue
                ws_f, ws_v = wb_f[name], wb_v[name]
                cells = []
                for rn, (row_f, row_v) in enumerate(zip(ws_f.iter_rows(), ws_v.iter_rows()),
                                                    start=1):
                    if rn > max_rows:
                        break
                    for cf, cv in zip(row_f, row_v):
                        if isinstance(cf.value, str) and cf.value.startswith('='):
                            cells.append((cf.coordinate, cv.value))
                if not cells:
                    out[name] = {'has_formula': False, 'formula_count': 0,
                                 'no_cache': False, 'no_cache_count': 0,
                                 'sample_locations': []}
                    continue
                no_cache = [coord for coord, cached in cells if cached is None]
                out[name] = {'has_formula': True, 'formula_count': len(cells),
                             'no_cache': bool(no_cache),
                             'no_cache_count': len(no_cache),
                             'sample_locations': no_cache[:5]}
            return out
        finally:
            wb_v.close()
    finally:
        wb_f.close()


def assert_no_uncalculated_formulas(path, sheet=None, what='操作'):
    """写类 CLI 的统一护栏：输入含「公式无缓存值」列时拒绝执行。

    pandas 会把这类公式列读成 NaN；若继续清洗/转换/转换，产出的是空列且
    公式被剥离，属于静默数据丢失。命中时 sys.exit 并给出 recalc 指引。
    """
    probe = probe_formulas(path, sheet=sheet)
    if not probe:
        return
    hit = {s: p for s, p in probe.items() if p.get('no_cache')}
    if not hit:
        return
    total = sum(p['no_cache_count'] for p in hit.values())
    sheets = '、'.join(hit.keys())
    sys.exit(
        f'错误：输入文件含 {total} 个无缓存值的公式单元格（抽查前 {PROBE_MAX_ROWS} 行，'
        f'sheet：{sheets}）。pandas 会把这些公式列读成空值，{what} 会剥离公式并产出空列。\n'
        f'请先处理公式：运行 recalc.py --cache（保公式 + 写缓存值，需 LibreOffice，'
        f'约 300MB 可在设置页安装）或 --bake（静态快照、公式会丢），确认数值正确后再重试。'
    )


def ensure_new_output(input_path, output_path):
    """输出不允许等于输入（防覆盖源文件，与 excel 角色约定一致）。"""
    norm = lambda p: os.path.normcase(os.path.abspath(p))
    if norm(input_path) == norm(output_path):
        sys.exit(f'错误：输出不能与输入是同一个文件（防覆盖），请换一个输出文件名 — {output_path}')
    return output_path


def save_frame(df, output_path, sheet_name='Sheet1'):
    """把 DataFrame 写出为 .xlsx/.csv/.tsv，自动建目录。

    - .csv 带 UTF-8 BOM（utf-8-sig，Excel 打开不乱码）
    - .xlsx 只写数据到指定 sheet 名（不保留样式/公式——需要保留原格式的场景不要走这里）
    """
    ext = os.path.splitext(output_path)[1].lower()
    if ext == '.xlsm':
        sys.exit('错误：pandas 写出 .xlsm 会丢失宏（实际只是换了后缀的无宏 xlsx）。'
                 '需要保留 VBA 时请用 openpyxl（load_workbook(keep_vba=True)）脚本处理，'
                 '或让用户在 Excel/WPS 中「另存为」；仅做纯数据清洗/转换时请输出 .xlsx')
    if ext not in ('.xlsx', '.csv', '.tsv'):
        sys.exit(f'错误：不支持的输出格式 "{ext}"，仅支持 .xlsx / .csv / .tsv')
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    if ext == '.csv':
        df.to_csv(output_path, index=False, encoding='utf-8-sig')
    elif ext == '.tsv':
        df.to_csv(output_path, sep='\t', index=False, encoding='utf-8-sig')
    else:
        df.to_excel(output_path, index=False, sheet_name=safe_sheet_name(sheet_name)[:31])
    return output_path, len(df)


def display_width(s):
    """单元格显示宽度：中日韩全角字符按 2 计（用于自动列宽）。"""
    import unicodedata
    return sum(2 if unicodedata.east_asian_width(ch) in 'FWA' else 1 for ch in str(s))


def format_pair_items(pairs_text, what='参数'):
    """解析 '键=值,键2=值2' 形式参数 → dict。

    兼容值内含 ASCII 逗号的情况（如数字格式 #,##0.00）——含 '=' 的片段开始新一项，
    不含 '=' 的片段视为上一项值的延续。键/值内不建议出现 '='。
    """
    items = {}
    text = str(pairs_text).strip()
    if not text:
        sys.exit(f'错误：{what} 为空')

    cur_key = None
    cur_val = None

    def flush():
        if cur_key is not None:
            items[cur_key] = cur_val

    for part in text.split(','):
        if '=' in part:
            flush()
            k, _, v = part.partition('=')
            cur_key, cur_val = k.strip(), v.strip()
        else:
            if cur_key is None:
                sys.exit(f'错误：{what} 片段 "{part}" 应为 键=值（用逗号分隔多项）')
            cur_val = (cur_val + ',' + part).strip()
    flush()
    if not items:
        sys.exit(f'错误：{what} 为空')
    return items


def split_csv(text, what='参数'):
    """逗号分隔 → 去空列表。"""
    parts = [p.strip() for p in str(text).split(',') if p and p.strip()]
    if not parts:
        sys.exit(f'错误：{what} 为空')
    return parts


def summarize_frame(df, sampled=False, note=None):
    """单 sheet/表的列统计 + 前 5 行预览（口径与 xl-read 模板一致）。

    - numeric 列：min/max/mean
    - object 列：唯一值数，唯一值 ≤ 10 时附 Top5 分布
    - 所有列：dtype / 非空 / 空值占比，sampled 标记是否基于采样
    """
    cols = []
    for c in df.columns:
        s = df[c]
        info = {
            'name': str(c),
            'dtype': str(s.dtype),
            'non_null': int(s.count()),
            'null_count': int(s.isnull().sum()),
            'null_pct': round(float(s.isnull().mean() * 100), 1),
            'sampled': sampled,
        }
        if pd.api.types.is_numeric_dtype(s):
            if s.count() > 0:
                info['min'] = float(s.min())
                info['max'] = float(s.max())
                info['mean'] = round(float(s.mean()), 2)
            else:
                info['min'] = info['max'] = info['mean'] = None
        elif pd.api.types.is_datetime64_any_dtype(s):
            # 日期列：给范围（min/max），回答「日期从哪天到哪天」
            valid = s.dropna()
            if len(valid):
                dmin, dmax = valid.min(), valid.max()
                to_d = lambda v: str(v.date()) if hasattr(v, 'date') else str(v)
                info['min'] = to_d(dmin)
                info['max'] = to_d(dmax)
                span = (dmax - dmin).days if hasattr(dmax - dmin, 'days') else None
                if span is not None:
                    info['span_days'] = int(span)
            else:
                info['min'] = info['max'] = None
        elif pd.api.types.is_string_dtype(s) or s.dtype == object:
            # 兼容 pandas 默认字符串 dtype（str）与传统 object dtype
            nunique = int(s.nunique())
            info['unique_count'] = nunique
            if nunique <= 10:
                info['top_values'] = {str(k): int(v) for k, v in s.value_counts().head(5).items()}
        cols.append(info)
    return {
        'columns': cols,
        'preview': df.head(5).fillna('').to_dict(orient='records'),
        'note': note,
    }


def to_markdown(df, limit=50):
    """手写 markdown 表格（不依赖 tabulate）。"""
    if limit and len(df) > limit:
        df = df.head(limit)

    def esc(v):
        if v is None:
            return ''
        return str(v).replace('|', '\\|').replace('\n', ' ')

    headers = [str(c) for c in df.columns]
    lines = [
        '| ' + ' | '.join(headers) + ' |',
        '| ' + ' | '.join('---' for _ in headers) + ' |',
    ]
    for _, row in df.iterrows():
        lines.append('| ' + ' | '.join(esc(v) for v in row) + ' |')
    return '\n'.join(lines)


def safe_sheet_name(name):
    """sheet 名 → 文件系统安全片段（用于多 sheet 展开输出文件名）。"""
    return ''.join('_' if ch in '\\/:*?"<>|' or ch.isspace() else ch for ch in str(name)) or 'sheet'
